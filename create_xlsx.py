import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

path = "E:\\vs-project\\health-check\\input\\Portfolio_Full_Health_QA_Hariharasuthan.xlsx"

tests = [
# ID, Area, Page/Element, Operation, Expected, Method, Priority
("NAV-001","Navigation","Home / HS.","Click logo","Returns to homepage/top hero without error","Manual browser test","High"),
("NAV-002","Navigation","About","Click ABOUT","Scrolls to About/Profile section","Manual browser test","High"),
("NAV-003","Navigation","Work","Click WORK","Scrolls to Selected Work section","Manual browser test","High"),
("NAV-004","Navigation","Experience","Click EXPERIENCE","Scrolls to Experience/Timeline section","Manual browser test","High"),
("NAV-005","Navigation","Contact","Click CONTACT","Scrolls to Contact section","Manual browser test","Critical"),
("NAV-006","Navigation","Scroll to Explore","Click/activate CTA","Moves user into portfolio content","Manual browser test","Medium"),
("NAV-007","Navigation","All anchor links","Click each internal anchor","Correct section opens; no wrong offset/blank area","Manual browser test","High"),
("NAV-008","Navigation","Mobile navigation","Open/close menu","Menu opens, closes, and links work","Mobile browser","Critical"),

("TAB-001","Page/Section","Hero","Load homepage","Hero renders completely; no broken assets/console-blocking errors","Manual + DevTools","Critical"),
("TAB-002","Page/Section","About / Profile","Open section","Profile text, image, education, email and focus are visible","Manual browser test","High"),
("TAB-003","Page/Section","Expertise / Toolkit","Review six expertise cards","All six cards render correctly and are readable","Manual browser test","High"),
("TAB-004","Page/Section","Selected Work","Review project list","All six project entries render and are clickable","Manual browser test","Critical"),
("TAB-005","Page/Section","Experience / Timeline","Review timeline","Adela, Nila and education entries render correctly","Manual browser test","High"),
("TAB-006","Page/Section","Contact","Review contact area","Contact email and social links are visible and usable","Manual browser test","Critical"),

("PROJ-001","Project","Image Retexturing","Open case study","IR_portfolio-details.html loads correctly","Browser/link test","Critical"),
("PROJ-002","Project","Image Retexturing","Click All Work","Returns to Selected Work/home section","Manual browser test","High"),
("PROJ-003","Project","Image Retexturing","Load project images","All four project images load without broken-image icons","Manual browser test","High"),
("PROJ-004","Project","Image Retexturing","Click View Source","Correct GitHub repository opens","External-link test","High"),
("PROJ-005","Project","Image Retexturing","Click Home","Returns to portfolio homepage","Manual browser test","High"),
("PROJ-006","Project","Brake Pad Inspection","Open case study","Correct case-study page loads","Link test","Critical"),
("PROJ-007","Project","Audio Diarization","Open case study","Correct case-study page loads","Link test","Critical"),
("PROJ-008","Project","R2+1D Fraud Detection","Open case study","Correct case-study page loads","Link test","Critical"),
("PROJ-009","Project","DOTfit Desk","Open case study","Correct case-study page loads","Link test","Critical"),
("PROJ-010","Project","Real-time Ball Tracking","Open case study","Correct case-study page loads","Link test","Critical"),

("EMAIL-001","Email/Contact","Email address","Click email on homepage","Mail client/appropriate email action opens with correct recipient","Manual browser test","Critical"),
("EMAIL-002","Email/Contact","Email address","Click email in About","Email link works and recipient is correct","Manual browser test","High"),
("EMAIL-003","Email/Contact","Email address","Click email in Contact","Email link works and recipient is correct","Manual browser test","Critical"),
("EMAIL-004","Email/Contact","Email operation","Send test email","Message is received at intended inbox","Real email test","Critical"),
("EMAIL-005","Email/Contact","Email operation","Reply to test email","Reply reaches test sender correctly","Real email test","High"),
("EMAIL-006","Email/Contact","Email operation","Check malformed/blocked email client behavior","User gets a usable fallback/contact path","Manual browser test","Medium"),

("SOC-001","Social","GitHub","Click GitHub","Correct GitHub profile opens","External-link test","Critical"),
("SOC-002","Social","LinkedIn","Click LinkedIn","Correct LinkedIn profile opens","External-link test","Critical"),
("SOC-003","Social","GitHub","Check profile accessibility","Profile is public and reachable without authentication","Browser test","High"),
("SOC-004","Social","External links","Open in new tab behavior","External navigation does not unexpectedly lose portfolio","Manual browser test","Medium"),

("CV-001","Resume/CV","Resume/CV link if present","Click Resume/CV","Correct current CV opens/downloads","Manual browser test","High"),
("CV-002","Resume/CV","Resume/CV","Open downloaded file","File opens and is readable","Manual file test","High"),
("CV-003","Resume/CV","Resume/CV","Check latest details","Experience/skills/contact details match portfolio","Content QA","High"),

("UI-001","UI/UX","Hero","Scroll through hero","No jump, clipping, or stuck animation","Manual browser test","Medium"),
("UI-002","UI/UX","Project cards","Hover/click cards","Hover states work and click target is obvious","Manual browser test","High"),
("UI-003","UI/UX","Experience cards","Scroll timeline","Animation/layout does not hide text","Manual browser test","Medium"),
("UI-004","UI/UX","Images","Inspect image scaling","No distortion or unexpected cropping","Manual browser test","Medium"),
("UI-005","UI/UX","Footer/contact","Scroll to bottom","Footer/contact content is reachable and not overlapped","Manual browser test","High"),
("UI-006","UI/UX","Back-to-top behavior","Scroll down then return","User can return to top using available navigation/logo","Manual browser test","Medium"),

("MOB-001","Responsive","Homepage","Test 320px width","No horizontal overflow; core content usable","Chrome DevTools","Critical"),
("MOB-002","Responsive","Homepage","Test 375px width","Layout and typography remain usable","Chrome DevTools","High"),
("MOB-003","Responsive","Homepage","Test 768px width","Tablet layout is stable","Chrome DevTools","High"),
("MOB-004","Responsive","Project pages","Test mobile project page","Images/text/buttons remain usable","Chrome DevTools","Critical"),
("MOB-005","Responsive","Contact","Test mobile contact area","Email/social links are tappable","Real device/DevTools","Critical"),
("MOB-006","Responsive","Navigation","Tap mobile navigation","No inaccessible or overlapping menu elements","Real device/DevTools","Critical"),

("BROWSER-001","Cross-browser","Homepage","Chrome desktop","All major operations work","Browser matrix","High"),
("BROWSER-002","Cross-browser","Homepage","Edge desktop","All major operations work","Browser matrix","Medium"),
("BROWSER-003","Cross-browser","Homepage","Firefox desktop","All major operations work","Browser matrix","Medium"),
("BROWSER-004","Cross-browser","Homepage","Safari/iOS if available","All major operations work","Browser matrix","Medium"),

("ASSET-001","Assets","Homepage","Check images","No 404/broken images","DevTools Network","High"),
("ASSET-002","Assets","Homepage","Check CSS","Stylesheets load successfully","DevTools Network","Critical"),
("ASSET-003","Assets","Homepage","Check JavaScript","Scripts load without fatal errors","DevTools Console/Network","Critical"),
("ASSET-004","Assets","Project pages","Check project assets","All project-specific CSS/JS/images load","DevTools Network","High"),
("ASSET-005","Assets","All pages","Check favicon","Favicon loads and appears in browser tab","Browser test","Low"),

("PERF-001","Performance","Homepage","Run Lighthouse mobile","Record Performance score and Core Web Vitals","Lighthouse","High"),
("PERF-002","Performance","Homepage","Run Lighthouse desktop","Record Performance score and Core Web Vitals","Lighthouse","Medium"),
("PERF-003","Performance","Homepage","Check LCP","Target <= 2.5s where measurable","Lighthouse/PageSpeed","High"),
("PERF-004","Performance","Homepage","Check INP","Target <= 200ms where measurable","PageSpeed/CrUX","High"),
("PERF-005","Performance","Homepage","Check CLS","Target <= 0.1 where measurable","Lighthouse/PageSpeed","High"),
("PERF-006","Performance","Homepage","Check oversized assets","Identify unnecessarily large images/fonts/scripts","Lighthouse/Network","Medium"),

("SEO-001","SEO","Homepage","Check title","Unique descriptive title exists","HTML inspection","High"),
("SEO-002","SEO","Homepage","Check meta description","Useful meta description exists","HTML inspection","Medium"),
("SEO-003","SEO","Homepage","Check H1","Primary H1 is present and meaningful","HTML inspection","High"),
("SEO-004","SEO","Homepage","Check canonical","Canonical points to intended portfolio URL","HTML inspection","High"),
("SEO-005","SEO","Site","Check robots.txt","Robots policy exists and does not block important pages","Browser","High"),
("SEO-006","SEO","Site","Check sitemap.xml","Sitemap exists if used and contains intended URLs","Browser","Medium"),
("SEO-007","SEO","Social sharing","Check OG metadata","Preview has correct title/image/description","HTML/social debugger","Medium"),

("GEO-001","GEO/AI Search","Homepage","Check identity clarity","Name, role and expertise are explicit","Manual review","High"),
("GEO-002","GEO/AI Search","Projects","Check project descriptions","Each project clearly states problem, approach, technology and outcome","Manual review","High"),
("GEO-003","GEO/AI Search","Experience","Check factual consistency","Dates, roles and skills are consistent","Manual review","High"),
("GEO-004","GEO/AI Search","Contact","Check discoverability","Professional contact path is explicit and consistent","Manual review","Medium"),

("A11Y-001","Accessibility","Homepage","Keyboard-only navigation","All interactive elements can be reached/used","Manual keyboard test","Critical"),
("A11Y-002","Accessibility","Homepage","Visible focus","Focus indicator is visible","Manual keyboard test","High"),
("A11Y-003","Accessibility","Images","Check alt text","Meaningful images have useful alt text","HTML/accessibility scan","High"),
("A11Y-004","Accessibility","Homepage","Check contrast","Text/controls meet contrast requirements","Lighthouse/axe","High"),
("A11Y-005","Accessibility","Animation","Reduced motion","Non-essential motion can be reduced/disabled","DevTools/manual","Medium"),

("SEC-001","Security","Source/repository","Search for secrets","No API keys, tokens, passwords or private data exposed","Repository/source scan","Critical"),
("SEC-002","Security","HTTPS","Open HTTPS","Site loads securely without certificate warning","Browser","Critical"),
("SEC-003","Security","Headers","Inspect response headers","Security headers are reviewed and intentional","Header scanner","Medium"),
("SEC-004","Security","Third-party assets","Inspect external resources","Only trusted external resources are loaded","Network/source review","Medium"),

("MON-001","Monitoring","Production site","Uptime check","Site is reachable from external monitor","Monitoring setup","High"),
("MON-002","Monitoring","Links","Broken-link scan","No broken internal/external links","Crawler","High"),
("MON-003","Monitoring","Email","Periodic email test","Contact/email path remains operational","Scheduled manual test","High"),
("MON-004","Monitoring","Deployment","Fresh deployment test","Latest GitHub Pages deployment matches intended commit","GitHub/deployment review","Medium"),
]

cols = ["Test ID","Area","Page / Element","Operation","Expected Result","Test Method","Priority",
        "Actual Result","Status","Severity if Failed","Evidence / Screenshot","Issue / Root Cause",
        "Recommended Fix","Retest Status","Tester","Test Date","Retest Date"]
df = pd.DataFrame(tests, columns=cols[:7])
for c in cols[7:]:
    df[c] = ""

summary = pd.DataFrame([
    ["Portfolio URL","https://hari7383.github.io/portfolio/"],
    ["Audit Date","2026-09-15"],
    ["Test Cases",len(df)],
    ["Passed",'=COUNTIF(QA!I:I,"Pass")'],
    ["Failed",'=COUNTIF(QA!I:I,"Fail")'],
    ["Not Tested",'=COUNTIF(QA!I:I,"Not Tested")'],
    ["N/A",'=COUNTIF(QA!I:I,"N/A")'],
    ["Health %",'=IFERROR(COUNTIF(QA!I:I,"Pass")/(COUNTIF(QA!I:I,"Pass")+COUNTIF(QA!I:I,"Fail")),0)'],
], columns=["Metric","Value"])

with pd.ExcelWriter(path, engine="openpyxl") as writer:
    summary.to_excel(writer, sheet_name="Dashboard", index=False)
    df.to_excel(writer, sheet_name="QA", index=False)

    for name, data in {
        "Navigation": df[df["Area"]=="Navigation"],
        "Email_Operations": df[df["Area"]=="Email/Contact"],
        "Projects_Links": df[df["Area"].isin(["Project","Social","Resume/CV"])],
        "UI_Responsive": df[df["Area"].isin(["UI/UX","Responsive","Cross-browser"])],
        "Performance": df[df["Area"]=="Performance"],
        "SEO_GEO": df[df["Area"].isin(["SEO","GEO/AI Search"])],
        "Accessibility": df[df["Area"]=="Accessibility"],
        "Security": df[df["Area"]=="Security"],
        "Assets_Monitoring": df[df["Area"].isin(["Assets","Monitoring"])],
    }.items():
        data.to_excel(writer, sheet_name=name, index=False)

wb = load_workbook(path)
header_fill = PatternFill("solid", fgColor="17365D")
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="D9E2F3")

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)

qa = wb["QA"]
widths = [14,20,24,36,48,24,14,38,16,18,30,36,42,16,18,14,14]
for i,w in enumerate(widths,1):
    qa.column_dimensions[chr(64+i) if i<=26 else "A"].width = w

status_dv = DataValidation(type="list", formula1='"Not Tested,Pass,Fail,N/A"', allow_blank=False)
severity_dv = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=True)
retest_dv = DataValidation(type="list", formula1='"Not Tested,Pass,Fail,N/A"', allow_blank=True)
qa.add_data_validation(status_dv); status_dv.add(f"I2:I{qa.max_row}")
qa.add_data_validation(severity_dv); severity_dv.add(f"J2:J{qa.max_row}")
qa.add_data_validation(retest_dv); retest_dv.add(f"N2:N{qa.max_row}")

# Conditional formatting with formula rules
qa.conditional_formatting.add(f"I2:I{qa.max_row}",
    FormulaRule(formula=['$I2="Fail"'], fill=PatternFill("solid", fgColor="FFC7CE")))
qa.conditional_formatting.add(f"I2:I{qa.max_row}",
    FormulaRule(formula=['$I2="Pass"'], fill=PatternFill("solid", fgColor="C6EFCE")))
qa.conditional_formatting.add(f"I2:I{qa.max_row}",
    FormulaRule(formula=['$I2="Not Tested"'], fill=PatternFill("solid", fgColor="FFF2CC")))

dash = wb["Dashboard"]
dash.column_dimensions["A"].width = 25
dash.column_dimensions["B"].width = 55
dash["B2"].hyperlink = "https://hari7383.github.io/portfolio/"
dash["B2"].style = "Hyperlink"
dash["B8"].number_format = "0.0%"
dash["A10"] = "How to use"
dash["B10"] = "Run each test manually/with the listed tool. Set Status. For every Fail, record evidence, root cause and fix. Retest after fixing."
dash["A10"].font = Font(bold=True)
dash["B10"].alignment = Alignment(wrap_text=True, vertical="top")
dash.row_dimensions[10].height = 55

# Add known current site inventory
inv = wb.create_sheet("Site_Inventory")
inventory = [
["Type","Item","Current evidence / URL","Expected test"],
["Homepage","Portfolio","https://hari7383.github.io/portfolio/","Loads successfully"],
["Section","ABOUT","Homepage #about / About navigation","Correct scroll/anchor"],
["Section","WORK","Homepage #work / Work navigation","Correct scroll/anchor"],
["Section","EXPERIENCE","Homepage #experience / Experience navigation","Correct scroll/anchor"],
["Section","CONTACT","Homepage #contact / Contact navigation","Correct scroll/anchor"],
["Project","Image Retexturing using AI","IR_portfolio-details.html","Case study loads"],
["Project","Brake Pad Inspection","Homepage project link","Case study loads"],
["Project","Audio Diarization","Homepage project link","Case study loads"],
["Project","R2+1D Fraud Detection","Homepage project link","Case study loads"],
["Project","DOTfit Desk","Homepage project link","Case study loads"],
["Project","Real-time Ball Tracking","Homepage project link","Case study loads"],
["Contact","Email","hari191203@gmail.com","Clickable and receives test email"],
["Social","GitHub","https://github.com/Hari7383","Opens correct profile"],
["Social","LinkedIn","Homepage LinkedIn link","Opens correct profile"],
]
for row in inventory:
    inv.append(row)
for cell in inv[1]:
    cell.fill = header_fill; cell.font = header_font
inv.column_dimensions["A"].width = 18
inv.column_dimensions["B"].width = 32
inv.column_dimensions["C"].width = 65
inv.column_dimensions["D"].width = 35
for row in inv.iter_rows():
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(path)
print(path)
